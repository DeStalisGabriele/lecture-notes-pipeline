"""Append Markdown notes to an existing Word document, matching its conventions.

Supported Markdown (what the agents are told to write):
  # Chapter / ## Section      numbered automatically ("31." / "31.2"), continuing the document
  ## ... before any "#"       continues the last chapter (e.g. 30.3 after an existing 30.2)
  ### Heading                 bold paragraph
  - item                      "List Paragraph"
  **bold**, *italic*
  | a | b |                   table
  ![caption](image.png)       picture 15 cm wide + italic caption
  MASTRINO: Cassa (attivo)    T-account (Italian accounting): merged title row, Dare | Avere,
  DARE: 420 — description     one entry per line
  AVERE: 2.000 — description
  EQ: Q = 4 \cdot 3 = 12 || calcolo le radici
                              one step of a calculation: LaTeX (no $) and an optional comment.
                              Consecutive EQ lines become a borderless table, a native Word
                              equation on the left and the comment in grey italics on the right.

Equations need the optional `latex2mathml` package and Office's MML2OMML.XSL stylesheet
(LaTeX -> MathML -> OMML). Without them the formula is written as plain text.

If the document ends with a summary chapter (e.g. "33. Schema riassuntivo delle formule"),
the new text goes before it and the summary is renumbered, so it stays last.
"""
from __future__ import annotations

import os
import re
import shutil
import time
from pathlib import Path

import docx
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}")
NUMBER = re.compile(r"\s*(\d+)\.")
EQUATION = re.compile(r"^\s*EQ:\s*(.+?)\s*(?:\|\|\s*(.*))?$")
# Office ships the MathML -> OMML stylesheet; it cannot be redistributed, so it is looked up here
OFFICE_XSL = [Path(os.environ.get(v, "")) / "Microsoft Office" / "root" / f"Office{n}" / "MML2OMML.XSL"
              for v in ("ProgramFiles", "ProgramFiles(x86)") for n in (16, 15)]


def find_mml2omml(configured: Path | None = None) -> Path | None:
    for p in ([configured] if configured else []) + OFFICE_XSL:
        if p and p.is_file():
            return p
    return None


class _Equations:
    """LaTeX -> native Word equation (OMML), or None when the tools are missing."""

    def __init__(self, xsl: Path | None):
        self.transform = None
        try:
            from latex2mathml.converter import convert
            if xsl:
                self.convert, self.transform = convert, etree.XSLT(etree.parse(str(xsl)))
        except ImportError:
            pass
        self.failed: list[str] = []

    def omml(self, latex: str):
        if not self.transform:
            return None
        try:
            return self.transform(etree.fromstring(self.convert(latex).encode())).getroot()
        except Exception:  # better plain text than a lost step
            self.failed.append(latex)
            return None


def _equation_paragraph(paragraph, omath) -> None:
    # alone in the paragraph and inside oMathPara, so Word shows it in display style (full-size fractions)
    block = etree.SubElement(paragraph._p, qn("m:oMathPara"))
    etree.SubElement(etree.SubElement(block, qn("m:oMathParaPr")), qn("m:jc")).set(qn("m:val"), "left")
    block.append(omath)


def _runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _style_name(p) -> str | None:
    return p.style.name if p.style is not None else None


def outline(path: Path) -> list[str]:
    """Chapter and section titles already in the document."""
    d = docx.Document(str(path))
    return [p.text for p in d.paragraphs if _style_name(p) in ("Heading 1", "Heading 2")]


def append_markdown(doc_path: Path, md_path: Path, backup_dir: Path | None = None,
                    summary_heading: str = r"schema riassuntivo", image_width_cm: float = 15,
                    mml2omml_xsl: Path | None = None) -> dict:
    if doc_path.exists():
        if backup_dir:
            backup_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(doc_path, backup_dir / f"{doc_path.stem}-backup-{time.strftime('%Y%m%d-%H%M%S')}.docx")
        d = docx.Document(str(doc_path))
    else:
        d = docx.Document()

    # Look styles up by object: by name, python-docx maps "Heading 1" to "heading 1" and misses
    # the styles of documents saved by some Word versions.
    styles = {s.name: s for s in d.styles}

    def par(text="", style="Normal"):
        return d.add_paragraph(text, style=styles.get(style))

    h1 = [p for p in d.paragraphs if _style_name(p) == "Heading 1"]
    anchor = h1[-1] if h1 and re.search(summary_heading, h1[-1].text, re.I) else None
    before = [p for p in h1 if anchor is None or p._p is not anchor._p]
    chapter = max((int(m.group(1)) for p in before if (m := NUMBER.match(p.text))), default=0)
    section = 0
    if before:
        inside = False
        for p in d.paragraphs:
            if p._p is before[-1]._p:
                inside = True
            elif anchor is not None and p._p is anchor._p:
                break
            elif inside and _style_name(p) == "Heading 2":
                section += 1

    body = d.element.body
    marker = (body[-2] if body[-1].tag.endswith("sectPr") else body[-1]) if len(body) >= 2 else None
    table: list[str] = []
    account: dict | None = None
    steps: list[tuple[str, str]] = []
    eq = _Equations(find_mml2omml(mml2omml_xsl))
    added = {"chapters": [], "pictures": 0, "tables": 0, "t_accounts": 0, "equations": 0}

    def flush_steps():
        t = d.add_table(rows=len(steps), cols=2)
        t.autofit = False
        for i, (latex, comment) in enumerate(steps):
            left, right = t.cell(i, 0), t.cell(i, 1)
            left.width, right.width = Cm(10.5), Cm(6)
            right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for cell in (left, right):  # some air between steps even when Normal has no spacing
                cell.paragraphs[0].paragraph_format.space_before = Pt(3)
                cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            if (omath := eq.omml(latex)) is not None:
                _equation_paragraph(left.paragraphs[0], omath)
                added["equations"] += 1
            else:
                left.paragraphs[0].add_run(latex)
                if not eq.transform:
                    added["equations_as_text"] = "latex2mathml or MML2OMML.XSL not found"
            if comment:
                run = right.paragraphs[0].add_run(comment)
                run.italic = True
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        steps.clear()
        par()

    def flush_table():
        rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table if not TABLE_RULE.match(r)]
        table.clear()
        if not rows:
            return
        cols = max(len(r) for r in rows)
        t = d.add_table(rows=len(rows), cols=cols)
        if "Table Grid" in styles:
            t.style = styles["Table Grid"]
        for i, row in enumerate(rows):
            for j in range(cols):
                cell = t.cell(i, j)
                _runs(cell.paragraphs[0], row[j] if j < len(row) else "")
        par()
        added["tables"] += 1

    def flush_account():
        nonlocal account
        acc, account = account, None
        t = d.add_table(rows=3, cols=2)
        if "Table Grid" in styles:
            t.style = styles["Table Grid"]
        head = t.cell(0, 0).merge(t.cell(0, 1))
        head.paragraphs[0].add_run(acc["name"]).bold = True
        for j, side in enumerate(("Dare", "Avere")):
            t.cell(1, j).paragraphs[0].add_run(side).bold = True
        for j, entries in enumerate((acc["dare"], acc["avere"])):
            cell = t.cell(2, j)
            for k, entry in enumerate(entries):
                _runs(cell.paragraphs[0] if k == 0 else cell.add_paragraph(), entry)
        par()
        added["t_accounts"] += 1

    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if m := EQUATION.match(line):
            steps.append((m.group(1), (m.group(2) or "").strip()))
            continue
        if steps:
            flush_steps()
        if m := re.match(r"^\s*MASTRINO:\s*(.+)$", line):
            if table:
                flush_table()
            if account:
                flush_account()
            account = {"name": m.group(1).strip(), "dare": [], "avere": []}
            continue
        if account and (m := re.match(r"^\s*(DARE|AVERE):\s*(.*)$", line)):
            account[m.group(1).lower()].append(m.group(2).strip())
            continue
        if account:
            flush_account()
        if line.strip().startswith("|"):
            table.append(line)
            continue
        if table:
            flush_table()
        if not line.strip():
            continue
        if line.startswith("# "):
            chapter += 1
            section = 0
            par(f"{chapter}. {line[2:].strip()}", "Heading 1")
            added["chapters"].append(f"{chapter}. {line[2:].strip()}")
        elif line.startswith("## "):
            section += 1
            par(f"{chapter}.{section} {line[3:].strip()}", "Heading 2")
            added["chapters"].append(f"{chapter}.{section} {line[3:].strip()}")
        elif line.startswith("### "):
            par().add_run(line[4:].strip()).bold = True
        elif m := re.match(r"!\[(.*?)\]\((.+?)\)", line.strip()):
            img = Path(m.group(2))
            img = img if img.is_absolute() else md_path.parent / img
            if img.exists():
                d.add_picture(str(img), width=Cm(image_width_cm))
                added["pictures"] += 1
                if m.group(1):
                    par().add_run(m.group(1)).italic = True
        elif re.match(r"\s*[-*] ", line):
            _runs(par(style="List Paragraph"), re.sub(r"^\s*[-*] ", "", line))
        else:
            _runs(par(), line.strip())
    if steps:
        flush_steps()
    if table:
        flush_table()
    if account:
        flush_account()

    if anchor is not None and marker is not None:
        new = []
        el = marker.getnext()
        while el is not None and not el.tag.endswith("sectPr"):
            new.append(el)
            el = el.getnext()
        for el in new:
            anchor._p.addprevious(el)
        if (old := NUMBER.match(anchor.text)):
            for run in anchor.runs:
                if old.group(1) in run.text:
                    run.text = run.text.replace(old.group(1), str(chapter + 1), 1)
                    break
        added["summary_moved_to"] = chapter + 1

    if eq.failed:
        added["equations_as_text"] = eq.failed
    d.save(str(doc_path))
    return added
