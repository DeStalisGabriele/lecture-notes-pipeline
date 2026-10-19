import docx
import pytest

from lecture_notes.docx_writer import append_markdown, find_mml2omml, outline

STEPS = ("# Esercizio\nCalcolo della produzione.\n"
         "EQ: Q = K^{0.5} \\cdot L^{0.5} || funzione di produzione\n"
         "EQ: Q = \\frac{16}{4} = 4 || sostituisco i valori\n"
         "Fine.\n")


def make_notes(path, with_summary=True):
    d = docx.Document()
    d.add_paragraph("29. Le imposte", style="Heading 1")
    d.add_paragraph("29.1 IVA", style="Heading 2")
    d.add_paragraph("30. Le interferenze fiscali", style="Heading 1")
    d.add_paragraph("30.1 Costi indeducibili", style="Heading 2")
    d.add_paragraph("30.2 Dividendi", style="Heading 2")
    d.add_paragraph("Testo del capitolo 30.")
    if with_summary:
        d.add_paragraph("31. Schema riassuntivo delle formule", style="Heading 1")
        d.add_paragraph("UAI = Ricavi - Costi")
    d.save(path)


def write_md(tmp_path, text):
    md = tmp_path / "notes.md"
    md.write_text(text, encoding="utf-8")
    return md


def test_new_text_goes_before_the_summary_and_continues_numbering(tmp_path):
    notes = tmp_path / "TEORIA.docx"
    make_notes(notes)
    md = write_md(tmp_path, "## Imposte differite\nTesto.\n# Leasing\n## Leasing operativo\nTesto.\n")
    added = append_markdown(notes, md)
    assert outline(notes) == [
        "29. Le imposte", "29.1 IVA",
        "30. Le interferenze fiscali", "30.1 Costi indeducibili", "30.2 Dividendi", "30.3 Imposte differite",
        "31. Leasing", "31.1 Leasing operativo",
        "32. Schema riassuntivo delle formule",
    ]
    assert added["summary_moved_to"] == 32
    texts = [p.text for p in docx.Document(notes).paragraphs]
    assert texts[-1] == "UAI = Ricavi - Costi"  # summary body still last


def test_without_summary_text_is_appended(tmp_path):
    notes = tmp_path / "TEORIA.docx"
    make_notes(notes, with_summary=False)
    append_markdown(notes, write_md(tmp_path, "# Obbligazioni\n- emissione alla pari\n"))
    assert outline(notes)[-1] == "31. Obbligazioni"


def test_t_account_layout(tmp_path):
    notes = tmp_path / "TEORIA.docx"
    md = write_md(tmp_path, "# Prova\nMASTRINO: Cassa (attivo)\nDARE: 420 — incasso dividendi\n"
                            "AVERE: 2.000 — acquisto partecipazione\nAVERE: 100 — spese\nFine.\n")
    added = append_markdown(notes, md)
    t = docx.Document(notes).tables[0]
    rows = [[c.text for c in r.cells] for r in t.rows]
    assert rows[0] == ["Cassa (attivo)", "Cassa (attivo)"]  # merged title row
    assert rows[1] == ["Dare", "Avere"]
    assert rows[2] == ["420 — incasso dividendi", "2.000 — acquisto partecipazione\n100 — spese"]
    assert added["t_accounts"] == 1


def test_markdown_table_bold_and_backup(tmp_path):
    notes = tmp_path / "TEORIA.docx"
    make_notes(notes, with_summary=False)
    md = write_md(tmp_path, "# Confronto\n| | Civilistico | Fiscale |\n|---|---|---|\n| **Imposte** | 1.200 | 800 |\n")
    append_markdown(notes, md, backup_dir=tmp_path / "bk")
    t = docx.Document(notes).tables[0]
    assert [c.text for c in t.rows[1].cells] == ["Imposte", "1.200", "800"]
    assert t.rows[1].cells[0].paragraphs[0].runs[0].bold
    assert len(list((tmp_path / "bk").glob("TEORIA-backup-*.docx"))) == 1


def test_equation_steps_without_stylesheet_fall_back_to_text(tmp_path, monkeypatch):
    monkeypatch.setattr("lecture_notes.docx_writer.OFFICE_XSL", [])
    notes = tmp_path / "TEORIA.docx"
    added = append_markdown(notes, write_md(tmp_path, STEPS))
    t = docx.Document(notes).tables[0]
    assert [[c.text for c in r.cells] for r in t.rows] == [
        [r"Q = K^{0.5} \cdot L^{0.5}", "funzione di produzione"],
        [r"Q = \frac{16}{4} = 4", "sostituisco i valori"],
    ]
    assert added["equations"] == 0 and "equations_as_text" in added
    assert docx.Document(notes).paragraphs[-1].text == "Fine."


@pytest.mark.skipif(find_mml2omml() is None, reason="needs Office's MML2OMML.XSL")
def test_equation_steps_become_word_equations(tmp_path):
    pytest.importorskip("latex2mathml")
    notes = tmp_path / "TEORIA.docx"
    added = append_markdown(notes, write_md(tmp_path, STEPS))
    assert added["equations"] == 2 and "equations_as_text" not in added
    t = docx.Document(notes).tables[0]
    m = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    assert len(t._tbl.findall(f".//{m}oMathPara")) == 2
    assert t.cell(1, 1).paragraphs[0].runs[0].italic
